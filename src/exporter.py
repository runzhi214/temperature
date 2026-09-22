from datetime import datetime
from typing import List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.calculator import AlertType, CalculationResult, format_duration
from src.logger import get_logger

HIGH_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
LOW_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=12, color="4472C4")
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
CENTER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)


def export_to_excel(results, filepath, range_labels=None):
    logger = get_logger()
    wb = Workbook()
    wb.remove(wb.active)

    for i, (result, label) in enumerate(_zip_results(results, range_labels)):
        sheet_name = f"区间{i + 1}" if label is None else label
        if len(sheet_name) > 31:
            sheet_name = sheet_name[:31]
        ws = wb.create_sheet(title=sheet_name)
        _write_sheet(ws, result, i + 1)

    try:
        wb.save(filepath)
        logger.info(f"Excel 导出完成: {filepath}")
    except PermissionError:
        logger.error("Excel 导出失败: 文件被占用, 请关闭后重试")
        raise
    except OSError as e:
        logger.error(f"Excel 导出失败: {e}")
        raise


def _zip_results(results, range_labels):
    if range_labels is None:
        return [(r, None) for r in results]
    return list(zip(results, range_labels))


def _write_sheet(ws, result, seq):
    row = 1
    ws.cell(row=row, column=1, value=f"温度监控数据分析报告 - 区间{seq}").font = TITLE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.cell(row=row, column=1).alignment = CENTER_ALIGN
    row += 2

    row = _write_summary_section(ws, row, result)
    row += 1
    row = _write_detail_section(ws, row, result)
    row += 1
    _write_statistics_section(ws, row, result)

    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 22


def _write_summary_section(ws, start_row, result):
    ws.cell(row=start_row, column=1, value="信息汇总").font = SECTION_FONT

    headers = ["序号", "状态", "持续时间", "时间占比", "报警次数"]
    row = start_row + 1
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
    row += 1

    for idx, s in enumerate(result.summary, 1):
        values = [idx, s.status, s.duration, f"{s.percentage}%", s.alert_count]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.alignment = CENTER_ALIGN
            cell.border = THIN_BORDER
            if "高温" in s.status:
                cell.fill = HIGH_FILL
            elif "低温" in s.status:
                cell.fill = LOW_FILL
        row += 1

    return row


def _write_detail_section(ws, start_row, result):
    alert_segments = [s for s in result.segments
                      if s.alert_type != AlertType.NORMAL]

    ws.cell(row=start_row, column=1, value="报警明细").font = SECTION_FONT

    headers = ["序号", "开始时间", "结束时间", "类型", "时长"]
    row = start_row + 1
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
    row += 1

    if not alert_segments:
        ws.cell(row=row, column=1, value="无报警记录").alignment = CENTER_ALIGN
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        row += 1
        return row

    for idx, seg in enumerate(alert_segments, 1):
        start_str = seg.start.strftime("%Y-%m-%d %H:%M")
        end_str = seg.end.strftime("%Y-%m-%d %H:%M")
        type_label = seg.alert_type.value
        duration_str = format_duration(seg.duration_seconds)
        values = [idx, start_str, end_str, type_label, duration_str]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.alignment = CENTER_ALIGN
            cell.border = THIN_BORDER
            if seg.alert_type == AlertType.HIGH:
                cell.fill = HIGH_FILL
            elif seg.alert_type == AlertType.LOW:
                cell.fill = LOW_FILL
        row += 1

    return row


def _write_statistics_section(ws, start_row, result):
    ws.cell(row=start_row, column=1, value="统计信息").font = SECTION_FONT
    row = start_row + 1

    stats = result.statistics
    lines = [
        ("开始时间", stats.start_time.strftime("%Y-%m-%d %H:%M:%S") if stats.data_count > 0 else "—"),
        ("结束时间", stats.end_time.strftime("%Y-%m-%d %H:%M:%S") if stats.data_count > 0 else "—"),
        ("累计时长", stats.total_duration),
        ("温度最大值", f"{stats.temp_max}℃"),
        ("温度最小值", f"{stats.temp_min}℃"),
        ("温度平均值", f"{stats.temp_avg}℃"),
        ("MKT (平均动力学温度)", f"{stats.mkt}℃"),
        ("数据条数", str(stats.data_count)),
    ]

    for label, value in lines:
        cell_label = ws.cell(row=row, column=1, value=label)
        cell_label.font = Font(bold=True)
        cell_label.alignment = Alignment(horizontal='right', vertical='center')
        cell_value = ws.cell(row=row, column=2, value=value)
        cell_value.alignment = Alignment(horizontal='left', vertical='center')
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        row += 1

    return row
