import csv
import re
from datetime import datetime
from typing import List, Tuple

from src.logger import get_logger


def parse_flexible_time(value):
    value = str(value).strip()
    if not value:
        return None
    for fmt in ['%Y/%m/%d %H:%M', '%Y-%m-%d %H:%M', '%Y/%m/%d %H:%M:%S',
                '%Y-%m-%d %H:%M:%S', '%Y/%m/%d %H:%M:%S']:
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(value)
    except (ValueError, TypeError):
        pass
    m = re.match(r'(\d{4})/(\d{1,2})/(\d{1,2})\s+(\d{1,2}):(\d{2})', value)
    if m:
        return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)),
                        int(m.group(4)), int(m.group(5)))
    return None


def parse_range_file(filepath):
    logger = get_logger()
    ranges = []

    if filepath.lower().endswith('.csv'):
        encodings = ['gbk', 'gb18030', 'utf-8-sig', 'utf-8']
        rows = None
        for enc in encodings:
            try:
                with open(filepath, 'r', encoding=enc) as f:
                    rows = list(csv.reader(f))
                break
            except (UnicodeDecodeError, LookupError):
                continue
        if rows is None:
            raise ValueError("无法解码文件")
    else:
        from openpyxl import load_workbook
        wb = load_workbook(filepath, read_only=True)
        ws = wb.active
        rows = [[str(c.value) if c.value is not None else '' for c in row] for row in ws.iter_rows()]
        wb.close()

    for i, row in enumerate(rows):
        if i == 0:
            stripped = [str(c).strip().lower() for c in row]
            if any(k in s for s in stripped for k in ['序号', '开始', '结束', '时间', 'start', 'end', 'no']):
                continue
        if len(row) < 3:
            continue
        try:
            start = parse_flexible_time(row[1])
            end = parse_flexible_time(row[2])
            if start and end and start < end:
                ranges.append((start, end))
            else:
                logger.debug(f"第 {i+1} 行时间无效，跳过")
        except Exception:
            logger.debug(f"第 {i+1} 行解析失败，跳过")
            continue

    return ranges
