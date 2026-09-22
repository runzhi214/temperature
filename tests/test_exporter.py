import unittest
import os
import tempfile
from datetime import datetime

from openpyxl import load_workbook

from src.calculator import (
    AlertType, AlertSegment, CalculationResult,
    SummaryRow, StatisticsInfo
)
from src.exporter import export_to_excel


def make_result():
    segments = [
        AlertSegment(
            start=datetime(2024, 10, 1, 10, 0),
            end=datetime(2024, 10, 1, 10, 40),
            alert_type=AlertType.NORMAL,
            duration_minutes=40
        ),
        AlertSegment(
            start=datetime(2024, 10, 1, 10, 40),
            end=datetime(2024, 10, 1, 10, 52),
            alert_type=AlertType.HIGH,
            duration_minutes=12
        ),
        AlertSegment(
            start=datetime(2024, 10, 1, 10, 52),
            end=datetime(2024, 10, 1, 11, 0),
            alert_type=AlertType.NORMAL,
            duration_minutes=8
        ),
    ]
    summary = [
        SummaryRow("温度正常", "48分钟", 48, 80.0, "—"),
        SummaryRow("高温报警", "12分钟", 12, 20.0, 1),
        SummaryRow("低温报警", "0分钟", 0, 0.0, 0),
    ]
    stats = StatisticsInfo(
        start_time=datetime(2024, 10, 1, 10, 0),
        end_time=datetime(2024, 10, 1, 11, 0),
        total_duration="1小时0分钟",
        total_minutes=60,
        temp_max=35.0,
        temp_min=20.0,
        temp_avg=25.0,
        mkt=26.5,
        data_count=4
    )
    return CalculationResult(segments=segments, summary=summary, statistics=stats)


class TestExportToExcel(unittest.TestCase):

    def setUp(self):
        self.result = make_result()

    def _get_temp_path(self):
        fd, path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        return path

    def test_export_creates_file(self):
        path = self._get_temp_path()
        try:
            export_to_excel([self.result], path)
            self.assertTrue(os.path.exists(path))
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_export_multiple_sheets(self):
        path = self._get_temp_path()
        try:
            export_to_excel([self.result, self.result], path,
                            range_labels=["区间A", "区间B"])
            wb = load_workbook(path)
            self.assertEqual(len(wb.sheetnames), 2)
            self.assertIn("区间A", wb.sheetnames)
            self.assertIn("区间B", wb.sheetnames)
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_sheet_content_summary(self):
        path = self._get_temp_path()
        try:
            export_to_excel([self.result], path)
            wb = load_workbook(path)
            ws = wb[wb.sheetnames[0]]
            title = ws.cell(row=1, column=1).value
            self.assertIn("温度监控数据分析报告", title)
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_sheet_has_summary_headers(self):
        path = self._get_temp_path()
        try:
            export_to_excel([self.result], path)
            wb = load_workbook(path)
            ws = wb[wb.sheetnames[0]]
            found = False
            for row in ws.iter_rows(min_row=1, max_row=10, max_col=5):
                values = [c.value for c in row]
                if "序号" in values and "状态" in values and "持续时间" in values:
                    found = True
                    break
            self.assertTrue(found)
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_sheet_has_detail_section(self):
        path = self._get_temp_path()
        try:
            export_to_excel([self.result], path)
            wb = load_workbook(path)
            ws = wb[wb.sheetnames[0]]
            found = False
            for row in ws.iter_rows(min_row=1, max_row=30, max_col=5):
                values = [c.value for c in row if c.value]
                for v in values:
                    if v == "报警明细":
                        found = True
                        break
            self.assertTrue(found)
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_sheet_has_statistics_section(self):
        path = self._get_temp_path()
        try:
            export_to_excel([self.result], path)
            wb = load_workbook(path)
            ws = wb[wb.sheetnames[0]]
            found = False
            for row in ws.iter_rows(min_row=1, max_row=40, max_col=5):
                values = [c.value for c in row if c.value]
                for v in values:
                    if v and "MKT" in str(v):
                        found = True
                        break
            self.assertTrue(found)
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_detail_has_alert_segments(self):
        path = self._get_temp_path()
        try:
            export_to_excel([self.result], path)
            wb = load_workbook(path)
            ws = wb[wb.sheetnames[0]]
            high_found = False
            for row in ws.iter_rows(min_row=1, max_row=30, max_col=5):
                values = [c.value for c in row]
                if "高温报警" in values:
                    high_found = True
                    break
            self.assertTrue(high_found)
        finally:
            if os.path.exists(path):
                os.unlink(path)


if __name__ == '__main__':
    unittest.main()
